"""Google Classroom materials from worksheets.json and guide.json.
Writes classroom/sheets/NN.html (one Google Doc each, uploaded through Drive),
classroom/teacher.html and classroom/Нұсқаулық.docx.
Usage: python3 -I classroom.py worksheets.json guide.json outdir site_url"""
import html, json, os, re, sys

wsj, gj, outdir, site = sys.argv[1:5]
WB = json.load(open(wsj, encoding='utf8'))
GUIDE = json.load(open(gj, encoding='utf8'))['sections']
sheets = sorted(WB['worksheets'], key=lambda w: w['order'])
os.makedirs(os.path.join(outdir, 'sheets'), exist_ok=True)
E = html.escape

GENRE = {'poetry': 'Поэзия', 'prose': 'Проза', 'drama': 'Драма', 'children': 'Балалар әдебиеті', 'folklore': 'Фольклор'}
TB = 'border="1" cellpadding="6" style="border-collapse:collapse;width:100%"'

def box(lines):
    return '<table %s><tr><td>%s</td></tr></table>' % (TB, '<br>' * (lines - 1) + '&nbsp;')

def field_html(f):
    lab = '<p><b>%s</b></p>' % E(f['label'])
    t = f['type']
    if t == 'text':
        return lab + box(1)
    if t == 'long':
        return lab + box(4)
    if t == 'list':
        rows = ''.join('<tr><td style="width:6%%">%d)</td><td>&nbsp;</td></tr>' % i for i in (1, 2, 3))
        return lab + '<table %s>%s</table>' % (TB, rows)
    if t in ('dropdown', 'tags'):
        opts = '&nbsp;&nbsp;&nbsp; '.join('☐ ' + E(o) for o in f['options'])
        return lab + '<p>%s</p>' % opts + (box(2) if t == 'tags' else '')
    if t == 'table':
        head = ''.join('<td><b>%s</b></td>' % E(c) for c in f['columns'])
        body = ''.join('<tr>%s</tr>' % ''.join('<td><br>&nbsp;</td>' for _ in f['columns']) for _ in range(f['rows']))
        return lab + '<table %s><tr>%s</tr>%s</table>' % (TB, head, body)
    if t == 'rubric':
        levels = sorted(f['criteria'][0]['levels'], key=int, reverse=True)
        names = {'3': 'жоғары', '2': 'орта', '1': 'бастапқы'}
        head = '<td><b>Критерий</b></td>' + ''.join('<td><b>%s балл: %s</b></td>' % (l, names.get(l, '')) for l in levels) + '<td><b>Балл</b></td>'
        body = ''.join('<tr><td><b>%s</b></td>%s<td>&nbsp;</td></tr>' % (E(c['name']), ''.join('<td>%s</td>' % E(c['levels'][l]) for l in levels)) for c in f['criteria'])
        bands = '; '.join('%d–%d балл: %s' % (b['min'], b['max'], b['label']) for b in f['bands'])
        return lab + '<table %s><tr>%s</tr>%s</table><p><i>%s</i></p>' % (TB, head, body, E(bands))
    raise SystemExit('unknown field type ' + t)

def refs(g):
    return re.findall(r'\d+(?:\.\d+)?', g or '')

def sheet_title(w):
    return '%02d %s' % (w['order'], w['title'])

def sheet_html(w):
    parts = ['<h1>%d-парақ. %s</h1>' % (w['order'], E(w['title']))]
    if w.get('audience') == 'teacher':
        parts.append('<table %s><tr><td>Шығарма:</td><td>Автор:</td><td>Сынып:</td></tr></table>' % TB)
    else:
        parts.append('<table %s><tr><td>Оқушының аты-жөні:</td><td>Сынып:</td><td>Күні:</td></tr>'
                     '<tr><td colspan="2">Шығарма:</td><td>Автор:</td></tr></table>' % TB)
    if refs(w.get('guide')):   # e.g. 'ҚМЖ шаблоны' names no guide section
        parts.append('<p><i>Нұсқаулықтың %s-тармағымен бірге қолданыңыз.</i></p>' % E(w['guide']))
    if w.get('note'):
        parts.append('<p><i>%s</i></p>' % E(w['note']))
    parts += [field_html(f) for f in w['fields']]
    if w.get('standard'):
        parts.append('<p><i>Бағдарламалық сәйкестік: %s</i></p>' % E(w['standard']))
    return '<html><body>%s</body></html>' % ''.join(parts)

for w in sheets:
    open(os.path.join(outdir, 'sheets', '%02d.html' % w['order']), 'w', encoding='utf8').write(sheet_html(w))
json.dump([{'order': w['order'], 'title': sheet_title(w), 'file': '%02d.html' % w['order']} for w in sheets],
          open(os.path.join(outdir, 'sheets', 'index.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)

# ------------------------------------------------------------ teacher doc
def sec_key(x):
    return tuple(int(p) for p in x.split('.'))

# Lesson split from guide chapter 16 (2.1–2.24 over seven lessons).
LESSONS = [('2.1–2.4', ['2.1', '2.2', '2.3', '2.4']), ('2.5–2.8', ['2.5', '2.6', '2.7', '2.8']),
           ('2.9–2.12', ['2.9', '2.10', '2.11', '2.12']), ('2.13–2.14, 2.24', ['2.13', '2.14', '2.24']),
           ('2.15–2.17, 2.23', ['2.15', '2.16', '2.17', '2.23']), ('2.18–2.19', ['2.18', '2.19']),
           ('2.20–2.22', ['2.20', '2.21', '2.22'])]
rubric = next(f for f in next(w for w in sheets if w['order'] == 31)['fields'] if f['type'] == 'rubric')

T = []
T.append('<h1>Көркем шығарманы талдау: Google Classroom-да қолдану</h1>')
T.append('<p>Бұл қалтада «Көркем шығарманы талдау» жұмыс дәптерінің 31 парағы бар. Әр парақ жеке Google Docs құжаты, сондықтан оны Classroom-да әр оқушыға жеке көшірме етіп беруге болады. Оқушы жауабын жиектелген орындарға жазады.</p>')
T.append('<h2>Тапсырма беру</h2><ol>'
         '<li>Classroom-да сыныпты ашып, <b>Тапсырмалар (Classwork)</b> бөліміне өтіңіз.</li>'
         '<li><b>Жасау (Create)</b>, содан кейін <b>Тапсырма (Assignment)</b> таңдаңыз.</li>'
         '<li>Тақырыбын жазыңыз, мысалы: «Ер Төстік. Сюжет картасы».</li>'
         '<li><b>Қосу (Add)</b>, содан кейін <b>Google Drive</b> басып, осы қалтадан керекті парақты таңдаңыз.</li>'
         '<li>Құжаттың жанындағы мәзірден <b>Әр оқушыға көшірме жасау (Make a copy for each student)</b> таңдаңыз. Бұл міндетті: әйтпесе оқушылар бір құжатқа бірге жазады.</li>'
         '<li>Мерзімін, тақырыпты (Topic) және рубриканы қосып, <b>Тапсыру (Assign)</b> басыңыз.</li></ol>')
T.append('<p>Оқушылар жауаптарын тапсырған соң, әр оқушының көшірмесін ашып, рубрика бойынша бағалап, пікір жаза аласыз.</p>')
T.append('<h2>Рубриканы Classroom-ға енгізу</h2>')
T.append('<p>Тапсырма ішінде <b>Рубрика (Rubric)</b>, содан кейін <b>Рубрика жасау (Create rubric)</b> басыңыз. Төмендегі 4 критерийді енгізіңіз. Әр критерийде 3 деңгей болады: 3, 2 және 1 балл. Бір рет жасаған рубриканы келесі тапсырмаларда <b>Қайта пайдалану (Reuse rubric)</b> арқылы таңдай аласыз.</p>')
levels = sorted(rubric['criteria'][0]['levels'], key=int, reverse=True)
names = {'3': 'Жоғары', '2': 'Орта', '1': 'Бастапқы'}
rows = ''.join('<tr><td><b>%s</b></td>%s</tr>' % (E(c['name']), ''.join('<td>%s</td>' % E(c['levels'][l]) for l in levels)) for c in rubric['criteria'])
T.append('<table %s><tr><td><b>Критерий</b></td>%s</tr>%s</table>' % (TB, ''.join('<td><b>%s балл: %s</b></td>' % (l, names[l]) for l in levels), rows))
T.append('<p>Жалпы балл 12. ' + E('; '.join('%d–%d: %s' % (b['min'], b['max'], b['label']) for b in rubric['bands'])) + '.</p>')
T.append('<h2>Қай парақты қашан беру керек</h2>')
T.append('<p>Нұсқаулықтың 16-бөлімі 2-бөлімді жеті сабаққа бөлуді ұсынады. Сол бөлімдерге сілтейтін парақтар:</p><table %s><tr><td><b>Сабақ</b></td><td><b>Нұсқаулық</b></td><td><b>Парақтар</b></td></tr>' % TB)
for i, (label, secs) in enumerate(LESSONS, 1):
    hit = [w for w in sheets if any(r in secs for r in refs(w['guide']))]
    T.append('<tr><td>%d</td><td>%s</td><td>%s</td></tr>' % (i, label, E('; '.join('%d. %s' % (w['order'], w['title']) for w in hit))))
T.append('</table>')
T.append('<p>Жалпы парақтар кез келген жанрға жарайды. Жанрға арналған парақтар: ' +
         E('; '.join('%s: %d-парақ' % (GENRE[g], w['order']) for w in sheets for g in w['genres'] if g in GENRE)) + '.</p>')
T.append('<p>Мұғалімге арналған парақтар: 24 (Шығарманы сабаққа бейімдеу) және 25 (Сабақ жоспары, ҚМЖ). Оларды оқушыларға бермеңіз. 26-парақ (Өзін-өзі бағалау) оқушыға арналған.</p>')
T.append('<h2>Нұсқаулық</h2>')
T.append('<p>Әдістемелік нұсқаулықтың толық мәтіні, іздеу мүмкіндігімен, сайтта тұр: <a href="%s">%s</a> («Нұсқаулық» қойындысы). Сол сайтта мұғалім өз талдауын жасап, 45 минуттық сабақ жоспарын шығара алады.</p>' % (E(site), E(site)))
T.append('<p>Нұсқаулықтың Word нұсқасы да бар. Оны Drive-қа жүктеп, Google Docs етіп ашсаңыз, бөлімдер бойынша мазмұн автоматты түрде құрылады.</p>')
T.append('<h2>Оқушы деректері</h2>')
T.append('<p>Оқушылардың жауаптары мектептің Google тіркелгісіндегі олардың өз көшірмелерінде сақталады. Сайт оқушылар туралы ешқандай дерек сақтамайды.</p>')
open(os.path.join(outdir, 'teacher.html'), 'w', encoding='utf8').write('<html><body>%s</body></html>' % ''.join(T))

# ------------------------------------------------------------ guide .docx
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
doc = Document()
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(11)
for st in ('Heading 1', 'Heading 2'):
    doc.styles[st].paragraph_format.keep_with_next = True
doc.add_heading('Көркем шығарманы талдау. Мұғалімге арналған әдістемелік нұсқаулық', 0)
doc.add_paragraph('Құрастырған: Бекжанова Сара. Алматы, 2026.')
for sec in GUIDE:
    if sec['level'] == 0:
        doc.add_heading('Алғы сөз және кіріспе', 1)
    elif sec['level'] == 1:
        doc.add_heading('%s-БӨЛІМ. %s' % (sec['id'], sec['title']), 1)
    else:
        doc.add_heading('%s. %s' % (sec['id'], sec['title']), 2)
    for p in sec['text'].split('\n\n'):
        if '\n' in p:                      # table kept as laid out in the PDF
            para = doc.add_paragraph()
            for i, line in enumerate(p.split('\n')):
                run = para.add_run(line)
                run.font.name = 'Courier New'
                run.font.size = Pt(7.5)
                if i < p.count('\n'):
                    run.add_break(WD_BREAK.LINE)
        elif re.match(r'^[•-] ', p):
            doc.add_paragraph(p[2:], style='List Bullet')
        else:
            doc.add_paragraph(p)
doc.save(os.path.join(outdir, 'Нұсқаулық.docx'))

sizes = [os.path.getsize(os.path.join(outdir, 'sheets', '%02d.html' % w['order'])) for w in sheets]
print('sheets: %d files, %d–%d bytes, total %d KB' % (len(sizes), min(sizes), max(sizes), sum(sizes) // 1024))
print('teacher.html:', os.path.getsize(os.path.join(outdir, 'teacher.html')), 'bytes')
print('guide docx:', os.path.getsize(os.path.join(outdir, 'Нұсқаулық.docx')) // 1024, 'KB')
